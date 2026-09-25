"""
Impact model for the recommended fix (route Category 23 straight to the group
that usually resolves it).

Inputs are of two kinds, kept visibly apart everywhere they appear:
  DATA        - measured from the incident log
  ASSUMPTION  - a judgement call, shown as a low / middle / high range

No fix has been deployed, so every output here is arithmetic on those inputs,
not a result. The same formulas are written to Excel so a stakeholder can
change the assumptions and see the effect.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PROCESSED = ROOT / "data" / "processed"
XLSX = ROOT / "Incident_Routing_Impact_Model.xlsx"
CATEGORY = "Category 23"
SCENARIOS = ["Low", "Middle", "High"]


def data_inputs() -> dict:
    d = pd.read_csv(PROCESSED / "incidents.csv", parse_dates=["opened_at"])
    c = d[(d.category == CATEGORY) & d.initial_group.notna() & d.final_group.notna()]
    metrics = pd.read_csv(PROCESSED / "current_state_metrics.csv", index_col="metric")["value"]
    policy = pd.read_csv(PROCESSED / "selective_policy.csv", index_col="metric")["value"]

    rule_group = policy["single_group_rule"].split("-> ")[1]
    core = c[c.opened_at.between(metrics["core_period_start"], metrics["core_period_end"] + " 23:59:59")]
    core = core.assign(month=core.opened_at.dt.to_period("M"), today_ok=core.initial_group == core.final_group,
                       rule_ok=core.final_group == rule_group, bounced=core.group_changes > 0)
    per_month = core.groupby("month").agg(
        incidents=("number", "size"), today_first_time_right=("today_ok", "mean"),
        resolved_by_rule_group=("rule_ok", "mean"), hand_off_rate=("bounced", "mean"))
    per_month["gain_points"] = per_month.resolved_by_rule_group - per_month.today_first_time_right
    per_month.round(4).to_csv(PROCESSED / "category23_by_month.csv", index_label="month")

    whole = core.rule_ok.mean() - core.today_ok.mean()
    holdout = float(policy["rule_accuracy_in_chosen_single_group"]) - float(policy["today_accuracy_in_chosen"])
    return {
        "rule": f"{CATEGORY} -> {rule_group}",
        "monthly_incidents": ", ".join(str(n) for n in per_month.incidents),
        "incidents_per_month": float(per_month.incidents.mean()),
        "gain_low": float(whole),
        "gain_high": holdout,
        "gain_middle": float((whole + holdout) / 2),
        "handoffs_per_bounced": float(c[c.group_changes > 0].group_changes.mean()),
        "extra_hours_like_for_like": float(metrics["like_for_like_extra_median_hours"]),
        "sla_gap": float(metrics["like_for_like_sla_gap_pts"]) / 100,
    }


def scenario_table(x: dict) -> pd.DataFrame:
    """Every input per scenario, and the outputs computed from them."""
    inputs = pd.DataFrame(
        {
            "incidents_per_month": [x["incidents_per_month"]] * 3,
            "first_time_right_gain": [x["gain_low"], x["gain_middle"], x["gain_high"]],
            "handoffs_avoided_per_incident": [1.0, 1.5, round(x["handoffs_per_bounced"], 1)],
            "like_for_like_extra_hours": [x["extra_hours_like_for_like"]] * 3,
            "share_of_gap_caused_by_handoff": [0.25, 0.50, 1.00],
            "sla_gap": [x["sla_gap"]] * 3,
            "minutes_per_handoff": [10, 15, 20],
            "cost_per_analyst_hour": [35, 45, 55],
            "one_off_effort_hours": [8, 8, 8],
        },
        index=SCENARIOS,
    )
    o = pd.DataFrame(index=SCENARIOS)
    o["incidents_routed_right_per_month"] = inputs.incidents_per_month * inputs.first_time_right_gain
    o["handoffs_avoided_per_month"] = o.incidents_routed_right_per_month * inputs.handoffs_avoided_per_incident
    o["elapsed_hours_saved_per_month"] = (
        o.incidents_routed_right_per_month * inputs.like_for_like_extra_hours * inputs.share_of_gap_caused_by_handoff)
    o["incidents_newly_meeting_sla_per_month"] = (
        o.incidents_routed_right_per_month * inputs.sla_gap * inputs.share_of_gap_caused_by_handoff)
    o["analyst_hours_saved_per_month"] = o.handoffs_avoided_per_month * inputs.minutes_per_handoff / 60
    o["analyst_cost_saved_per_month"] = o.analyst_hours_saved_per_month * inputs.cost_per_analyst_hour
    o["analyst_cost_saved_per_year"] = o.analyst_cost_saved_per_month * 12
    o["one_off_cost"] = inputs.one_off_effort_hours * inputs.cost_per_analyst_hour
    o["payback_days"] = o.one_off_cost / o.analyst_cost_saved_per_month * 30
    return pd.concat([inputs.add_prefix("in_"), o], axis=1)


def write_excel(x: dict, table: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    blue, bold = Font(color="0000FF"), Font(bold=True)
    shade = PatternFill("solid", fgColor="E3EFE9")
    warn = PatternFill("solid", fgColor="F4E6D6")

    ws["A1"] = f"Impact model: routing {CATEGORY} straight to the group that usually resolves it"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Blue = an input you can change. Black = a formula. DATA inputs are measured from the incident "
                "log; ASSUMPTION inputs are judgement calls. No fix has been deployed, so outputs are estimates, "
                "not results.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 44

    for col, head in enumerate(["", "Unit", *SCENARIOS, "Basis"], start=1):
        c = ws.cell(row=4, column=col, value=head)
        c.font, c.fill = bold, shade

    def cols(name):
        return list(table[name])

    rows = [  # (label, unit, values, number_format, kind, basis)
        (f"{CATEGORY} incidents per month", "incidents", cols("in_incidents_per_month"), "#,##0", "DATA",
         f"Average of the three core months ({x['monthly_incidents']})"),
        ("First-time-right gain from the rule", "points", cols("in_first_time_right_gain"), "0.0%", "DATA",
         "Low: whole Mar-May period. High: held-out latest data. Middle: midpoint"),
        ("Hand-offs avoided per incident routed right", "count", cols("in_handoffs_avoided_per_incident"), "0.0",
         "DATA / ASSUMPTION", f"High = observed average ({x['handoffs_per_bounced']:.1f}) among bounced incidents; Low and Middle are conservative"),
        ("Extra hours a bounced incident takes, like-for-like", "hours", cols("in_like_for_like_extra_hours"), "0.0",
         "DATA", "Median gap within the same resolving group (association, not proof of cause)"),
        ("Share of that gap assumed caused by the hand-off", "%", cols("in_share_of_gap_caused_by_handoff"), "0%",
         "ASSUMPTION", "Harder tickets get passed around, so only part of the gap is caused by the hand-off"),
        ("SLA attainment gap, like-for-like", "points", cols("in_sla_gap"), "0.0%", "DATA",
         "SLA met: no hand-off vs changed hands, within the same resolving group"),
        ("Analyst minutes per hand-off", "minutes", cols("in_minutes_per_handoff"), "0", "ASSUMPTION",
         "Time to read, triage and pass a ticket on. Not in the data; replace with your own"),
        ("Loaded cost per analyst hour", "$", cols("in_cost_per_analyst_hour"), "$#,##0", "ASSUMPTION",
         "Placeholder. Replace with the finance figure"),
        ("One-off effort to change the rule", "hours", cols("in_one_off_effort_hours"), "0", "ASSUMPTION",
         "A configuration change in the assignment rules"),
    ]
    first_input = 5
    for i, (label, unit, values, fmt, kind, basis) in enumerate(rows):
        r = first_input + i
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=unit)
        for j, v in enumerate(values):
            cell = ws.cell(row=r, column=3 + j, value=float(v))
            cell.font, cell.number_format = blue, fmt
            if kind == "ASSUMPTION":
                cell.fill = warn
        ws.cell(row=r, column=6, value=f"[{kind}] {basis}")

    out_start = first_input + len(rows) + 2  # header row for outputs
    for col, head in enumerate(["Outputs (formulas)", "Unit", *SCENARIOS, "How it is calculated"], start=1):
        c = ws.cell(row=out_start, column=col, value=head)
        c.font, c.fill = bold, shade

    R = {name: first_input + i for i, name in enumerate(
        ["inc", "gain", "hand", "hours", "share", "sla", "mins", "rate", "effort"])}
    outputs = [  # (label, unit, formula template using {c}, number_format, note)
        ("Incidents routed right that were not before", "per month", "={c}{inc}*{c}{gain}", "#,##0", "incidents x gain"),
        ("Hand-offs avoided", "per month", "={c}{o0}*{c}{hand}", "#,##0", "incidents routed right x hand-offs avoided"),
        ("Elapsed resolution hours saved", "per month", "={c}{o0}*{c}{hours}*{c}{share}", "#,##0",
         "incidents routed right x extra hours x share caused by the hand-off"),
        ("Incidents newly meeting SLA", "per month", "={c}{o0}*{c}{sla}*{c}{share}", "#,##0",
         "incidents routed right x SLA gap x share caused by the hand-off"),
        ("Analyst hours saved", "per month", "={c}{o1}*{c}{mins}/60", "#,##0", "hand-offs avoided x minutes / 60"),
        ("Analyst cost saved", "$ per month", "={c}{o4}*{c}{rate}", "$#,##0", "analyst hours x cost per hour"),
        ("Analyst cost saved", "$ per year", "={c}{o5}*12", "$#,##0", "assumes volume stays at the observed level"),
        ("One-off cost of the change", "$", "={c}{effort}*{c}{rate}", "$#,##0", "effort hours x cost per hour"),
        ("Payback", "days", "=IF({c}{o5}>0,{c}{o7}/{c}{o5}*30,\"n/a\")", "0.0", "one-off cost / monthly saving x 30"),
    ]
    o_rows = {f"o{i}": out_start + 1 + i for i in range(len(outputs))}
    for i, (label, unit, template, fmt, note) in enumerate(outputs):
        r = out_start + 1 + i
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=unit)
        for j in range(3):
            letter = get_column_letter(3 + j)
            # placeholders like {c}{inc} expand to a cell address such as C5
            cell = ws.cell(row=r, column=3 + j, value=template.format(c=letter, **R, **o_rows))
            cell.number_format = fmt
        ws.cell(row=r, column=6, value=note)

    for col, width in zip("ABCDEF", [50, 12, 13, 13, 13, 92]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C5"
    wb.save(XLSX)


def main():
    x = data_inputs()
    table = scenario_table(x)
    table.round(4).to_csv(PROCESSED / "impact_scenarios.csv", index_label="scenario")
    write_excel(x, table)
    pd.set_option("display.width", 200)
    print("rule:", x["rule"])
    print(table.filter(like="in_").round(3).T.to_string())
    print(table.drop(columns=table.filter(like="in_").columns).round(1).T.to_string())
    print("wrote", XLSX)


if __name__ == "__main__":
    main()
